from PyQt5.QtWidgets import QFileDialog, QLineEdit, QComboBox, QToolButton
import random
from openpyxl import load_workbook

def select_file(parent, line:QLineEdit, columns:QComboBox, generator:QToolButton, column_line:QLineEdit):
    file_path, _ = QFileDialog.getOpenFileName(
        parent,
        "Sélectionner un fichier Excel",
        "",
        "Fichiers Excel (*.xlsx *.xls)"
    )

    if file_path:
        line.setText(file_path)
        titles_row = find_header_row(file_path)
        c = get_headers(file_path, titles_row) # type: ignore
        columns.clear()
        columns.addItems(c)
        generator.setEnabled(True)
        columns.setVisible(True)
        column_line.setVisible(True)


def find_header_row(file_path, target="N° Place"):

    wb = load_workbook(file_path, read_only=True)
    ws = wb.active

    for row in ws.iter_rows(): # type: ignore
        values = [cell.value for cell in row]

        if target in values:
            for i in range(10):
                if row[i].value is not None:
                    return row[i].row
                else:
                    pass
    wb.close()
    return 1

def get_headers(file_path, header_row=2):

    wb = load_workbook(file_path, read_only=True)
    ws = wb.active

    headers = []

    for cell in ws[header_row]: # type: ignore
        if cell.value is not None:
            headers.append(cell.value)
    wb.close()
    return headers
"""
def create_random_place_column(
        input_file,
        source_column="N° Place",
        new_column="N° Place1"
):

    wb = load_workbook(input_file)

    ws = wb.active


    # Recherche de la ligne d'en-tête
    header_row = find_header_row(
        input_file,
        source_column
    )


    if header_row is None:
        raise Exception(
            f"Impossible de trouver {source_column}"
        )


    # Récupération des colonnes
    headers = {}

    for cell in ws[header_row]:

        if cell.value is not None:
            headers[cell.value] = cell.column



    if source_column not in headers:
        raise Exception(
            f"La colonne {source_column} n'existe pas"
        )


    source_col = headers[source_column]


    # Récupération des valeurs
    values = []

    for row in range(
        header_row + 1,
        ws.max_row + 1
    ):

        value = ws.cell(row, source_col).value

        if value is not None:
            values.append(value)



    if len(values) < 2:
        raise Exception(
            "Pas assez de valeurs pour mélanger"
        )



    # Création de la nouvelle colonne
    new_col = ws.max_column + 1


    ws.cell(
        header_row,
        new_col
    ).value = new_column



    # Mélange sans doublon de position
    shuffled = values.copy()


    tentatives = 0

    while True:

        random.shuffle(shuffled)

        if all(
            a != b
            for a, b in zip(values, shuffled)
        ):
            break


        tentatives += 1

        if tentatives > 1000:
            raise Exception(
                "Impossible de générer un mélange valide"
            )



    # Écriture des valeurs
    index = 0

    for row in range(
        header_row + 1,
        ws.max_row + 1
    ):

        if ws.cell(row, source_col).value is not None:

            ws.cell(
                row,
                new_col
            ).value = shuffled[index]

            index += 1



    # Sauvegarde directe
    wb.save(input_file)
"""

def create_random_place_column(
        input_file,
        source_column="N° Place",
        new_column="N° Place1",
        progress_callback=None
):

    wb = load_workbook(input_file)
    ws = wb.active


    # Recherche ligne d'en-tête
    header_row = find_header_row(
        input_file,
        source_column
    )


    if header_row is None:
        raise Exception(
            f"{source_column} introuvable"
        )


    headers = {}

    for cell in ws[header_row]: # type: ignore
        if cell.value is not None:
            headers[cell.value] = cell.column


    source_col = headers[source_column]


    # Récupération des valeurs
    values = []

    total_rows = ws.max_row - header_row # type: ignore


    for index, row in enumerate(
        range(header_row + 1, ws.max_row + 1),
        start=1
    ):

        value = ws.cell(row, source_col).value # type: ignore

        if value is not None:
            values.append(value)


        # Progression lecture
        if progress_callback:
            progress_callback(
                int(index / total_rows * 30)
            )


    # Mélange
    shuffled = values.copy()

    while True:

        random.shuffle(shuffled)

        if all(
            a != b
            for a, b in zip(values, shuffled)
        ):
            break


    if progress_callback:
        progress_callback(50)


    # Création colonne
    new_col = ws.max_column + 1 # type: ignore

    ws.cell(
        header_row,
        new_col
    ).value = new_column # type: ignore



    # Écriture
    total = len(values)

    index = 0

    for row in range(
        header_row + 1,
        ws.max_row + 1 # type: ignore
    ):

        if ws.cell(row, source_col).value is not None:# type: ignore

            ws.cell( 
                row,
                new_col
            ).value = shuffled[index] 

            index += 1


        if progress_callback:
            progress_callback(
                50 + int(index / total * 40)
            )


    if progress_callback:
        progress_callback(95)


    wb.save(input_file)
    wb.close()


    if progress_callback:
        progress_callback(100)

from PyQt5.QtCore import QObject, pyqtSignal


class ExcelWorker(QObject):

    progress = pyqtSignal(int)
    finished = pyqtSignal()
    error = pyqtSignal(str)


    def __init__(self, file_path, source_col, new_col):
        super().__init__()
        self.file_path = file_path
        self.source_col = source_col
        self.new_col = new_col


    def run(self):

        try:

            create_random_place_column(
                self.file_path,
                self.source_col,
                self.new_col,
                progress_callback=self.progress.emit
            )

            self.finished.emit()


        except Exception as e:
            self.error.emit(str(e))